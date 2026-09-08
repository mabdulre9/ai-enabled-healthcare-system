from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.config import CLINIC_NAME, CLINIC_ADDRESS, CLINIC_PHONE


def _set_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def _add_clinic_header(doc):
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(CLINIC_NAME)
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(102, 126, 234)

    header2 = doc.add_paragraph()
    header2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header2.add_run(f'{CLINIC_ADDRESS}\n{CLINIC_PHONE}')
    run.font.size = Pt(10)

    doc.add_paragraph()


def _add_footer(doc, text='CONFIDENTIAL MEDICAL RECORD\nFor authorized use only'):
    doc.add_paragraph()
    doc.add_paragraph('_' * 80)
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f'\nReport Generated: {datetime.now().strftime("%B %d, %Y at %I:%M %p")}\n')
    run.font.size = Pt(9)
    run.font.italic = True

    footer2 = doc.add_paragraph()
    footer2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer2.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(128, 128, 128)


def generate_visit_report_docx(patient, visit):
    doc = Document()
    _set_margins(doc)
    _add_clinic_header(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('VISIT REPORT')
    run.font.size = Pt(14)
    run.font.bold = True

    doc.add_paragraph('_' * 80)

    doc.add_heading('Patient Information', level=2)
    patient_info = doc.add_paragraph()
    patient_info.add_run("Name: ").bold = True
    patient_info.add_run(f"{patient.get('name', 'N/A')}\n")
    patient_info.add_run("Patient ID: ").bold = True
    patient_info.add_run(f"{patient.get('patientId', 'N/A')}\n")
    patient_info.add_run("Date of Birth: ").bold = True
    patient_info.add_run(f"{patient.get('dateOfBirth', 'N/A')} (Age: {patient.get('age', 'N/A')})\n")
    patient_info.add_run("Gender: ").bold = True
    patient_info.add_run(f"{patient.get('gender', 'N/A').title()}\n")

    doc.add_heading('Visit Information', level=2)
    visit_info = doc.add_paragraph()
    visit_info.add_run("Visit ID: ").bold = True
    visit_info.add_run(f"{visit.get('visitId', 'N/A')}\n")
    visit_info.add_run("Date: ").bold = True
    visit_info.add_run(f"{visit.get('date', 'N/A')}\n")
    visit_info.add_run("Type: ").bold = True
    visit_info.add_run(f"{visit.get('type', 'N/A')}\n")
    visit_info.add_run("Provider: ").bold = True
    visit_info.add_run(f"{visit.get('seenBy', 'N/A')}\n")

    doc.add_heading('Chief Complaint', level=2)
    doc.add_paragraph(visit.get('chiefComplaint', 'N/A'))

    if visit.get('vitals'):
        doc.add_heading('Vital Signs', level=2)
        vitals = visit['vitals']
        table = doc.add_table(rows=5, cols=2)
        table.style = 'Light Grid Accent 1'
        vitals_data = [
            ('Blood Pressure', vitals.get('bloodPressure', 'N/A')),
            ('Heart Rate', f"{vitals.get('heartRate', 'N/A')} bpm"),
            ('Temperature', f"{vitals.get('temperature', 'N/A')} F"),
            ('Weight', f"{vitals.get('weight', 'N/A')} lbs (BMI: {vitals.get('bmi', 'N/A')})"),
            ('Oxygen Saturation', f"{vitals.get('oxygenSat', 'N/A')}%")
        ]
        for i, (label, value) in enumerate(vitals_data):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
            table.rows[i].cells[1].text = value

    doc.add_heading('Symptoms', level=2)
    doc.add_paragraph(visit.get('symptoms', 'N/A'))

    doc.add_heading('Physical Examination', level=2)
    doc.add_paragraph(visit.get('examination', 'N/A'))

    doc.add_heading('Assessment & Diagnosis', level=2)
    doc.add_paragraph(visit.get('assessment', 'N/A'))

    doc.add_heading('Treatment Plan', level=2)
    doc.add_paragraph(visit.get('plan', 'N/A'))

    if visit.get('prescriptions') and len(visit['prescriptions']) > 0:
        doc.add_heading('Prescriptions', level=2)
        for rx in visit['prescriptions']:
            if rx.get('medication'):
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"{rx.get('medication')} ").bold = True
                p.add_run(f"{rx.get('dosage', '')} - {rx.get('frequency', '')}")
                if rx.get('quantity'):
                    p.add_run(f"\n   Quantity: {rx['quantity']}, Refills: {rx.get('refills', '0')}")

    if visit.get('labsOrdered'):
        doc.add_heading('Laboratory Tests Ordered', level=2)
        doc.add_paragraph(visit.get('labsOrdered'))

    if visit.get('followUp'):
        doc.add_heading('Follow-up', level=2)
        doc.add_paragraph(f"Scheduled for: {visit.get('followUp')}")

    if visit.get('notes'):
        doc.add_heading('Additional Notes', level=2)
        doc.add_paragraph(visit.get('notes'))

    _add_footer(doc)
    return doc


def generate_patient_report_docx(patient):
    doc = Document()
    _set_margins(doc)
    _add_clinic_header(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('COMPREHENSIVE MEDICAL RECORD')
    run.font.size = Pt(14)
    run.font.bold = True

    doc.add_paragraph('_' * 80)

    doc.add_heading('Patient Demographics', level=1)
    demo_table = doc.add_table(rows=6, cols=2)
    demo_table.style = 'Light Grid Accent 1'
    demo_data = [
        ('Full Name', patient.get('name', 'N/A')),
        ('Patient ID', patient.get('patientId', 'N/A')),
        ('Date of Birth', f"{patient.get('dateOfBirth', 'N/A')} (Age: {patient.get('age', 'N/A')})"),
        ('Gender', patient.get('gender', 'N/A').title()),
        ('Phone', patient.get('phone', 'N/A')),
        ('Email', patient.get('email', 'N/A'))
    ]
    for i, (label, value) in enumerate(demo_data):
        demo_table.rows[i].cells[0].text = label
        demo_table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
        demo_table.rows[i].cells[1].text = value

    if patient.get('address'):
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run('Address: ').bold = True
        p.add_run(patient.get('address'))

    if patient.get('emergencyContact'):
        ec = patient['emergencyContact']
        if ec.get('name'):
            doc.add_paragraph()
            p = doc.add_paragraph()
            p.add_run('Emergency Contact: ').bold = True
            p.add_run(f"{ec.get('name')} ({ec.get('relationship', 'N/A')}) - {ec.get('phone', 'N/A')}")

    doc.add_heading('Active Medical Conditions', level=1)
    if patient.get('activeConditions') and len(patient['activeConditions']) > 0:
        for cond in patient['activeConditions']:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{cond.get('condition', 'N/A')}").bold = True
            details = f"\n   Diagnosed: {cond.get('diagnosedDate', 'N/A')}"
            if cond.get('diagnosedBy'):
                details += f" by {cond['diagnosedBy']}"
            if cond.get('severity'):
                details += f" | Severity: {cond['severity']}"
            if cond.get('notes'):
                details += f"\n   Notes: {cond['notes']}"
            p.add_run(details)
    else:
        doc.add_paragraph('No active conditions recorded.')

    if patient.get('pastConditions') and len(patient['pastConditions']) > 0:
        doc.add_heading('Past Medical Conditions (Resolved)', level=2)
        for cond in patient['pastConditions']:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{cond.get('condition', 'N/A')}")
            p.add_run(f" ({cond.get('diagnosedDate', 'N/A')} to {cond.get('resolvedDate', 'N/A')})")

    doc.add_heading('Current Medications', level=1)
    if patient.get('activeMedications') and len(patient['activeMedications']) > 0:
        for med in patient['activeMedications']:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{med.get('medication', 'N/A')} ").bold = True
            p.add_run(f"{med.get('dosage', '')} - {med.get('frequency', '')} ({med.get('route', 'oral')})")
            details = f"\n   Started: {med.get('startedDate', 'N/A')}"
            if med.get('purpose'):
                details += f" | Purpose: {med['purpose']}"
            if med.get('prescribedBy'):
                details += f"\n   Prescribed by: {med['prescribedBy']}"
            p.add_run(details)
    else:
        doc.add_paragraph('No current medications.')

    doc.add_heading('ALLERGIES & ADVERSE REACTIONS', level=1)
    if patient.get('allergies') and len(patient['allergies']) > 0:
        allergy_table = doc.add_table(rows=len(patient['allergies']) + 1, cols=3)
        allergy_table.style = 'Medium Shading 1 Accent 2'
        header_cells = allergy_table.rows[0].cells
        header_cells[0].text = 'Allergen'
        header_cells[1].text = 'Reaction'
        header_cells[2].text = 'Severity'
        for i, cell in enumerate(header_cells):
            cell.paragraphs[0].runs[0].font.bold = True
        for i, allergy in enumerate(patient['allergies']):
            row = allergy_table.rows[i + 1]
            row.cells[0].text = allergy.get('allergen', 'N/A')
            row.cells[1].text = allergy.get('reaction', 'N/A')
            row.cells[2].text = allergy.get('severity', 'N/A').upper()
    else:
        doc.add_paragraph('No known allergies.')

    if patient.get('familyHistory') and len(patient['familyHistory']) > 0:
        doc.add_heading('Family Medical History', level=1)
        for family in patient['familyHistory']:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{family.get('relative', 'N/A')} ").bold = True
            p.add_run(f"({family.get('age', 'N/A')} years, {family.get('livingStatus', 'N/A')})")
            if family.get('conditions'):
                p.add_run('\n   Conditions: ' + family.get('conditions', ''))
            if family.get('notes'):
                p.add_run(f"\n   Notes: {family['notes']}")

    if patient.get('socialHistory'):
        doc.add_heading('Social History', level=1)
        sh = patient['socialHistory']
        social_table = doc.add_table(rows=5, cols=2)
        social_table.style = 'Light Grid Accent 1'
        social_data = [
            ('Smoking Status', sh.get('smokingStatus', 'N/A')),
            ('Alcohol Use', sh.get('alcoholUse', 'N/A')),
            ('Occupation', sh.get('occupation', 'N/A')),
            ('Exercise Habits', sh.get('exerciseHabits', 'N/A')),
            ('Diet', sh.get('diet', 'N/A'))
        ]
        for i, (label, value) in enumerate(social_data):
            social_table.rows[i].cells[0].text = label
            social_table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
            social_table.rows[i].cells[1].text = value

    if patient.get('visits') and len(patient['visits']) > 0:
        doc.add_heading('Visit History', level=1)
        doc.add_paragraph(f'Total Visits: {len(patient["visits"])}')
        recent_visits = patient['visits'][-5:]
        recent_visits.reverse()
        for visit in recent_visits:
            doc.add_paragraph()
            p = doc.add_paragraph()
            p.add_run(f"Visit Date: {visit.get('date', 'N/A')}").bold = True
            p.add_run(f" | Type: {visit.get('type', 'N/A')}")
            p2 = doc.add_paragraph(style='List Bullet 2')
            p2.add_run('Chief Complaint: ').bold = True
            p2.add_run(visit.get('chiefComplaint', 'N/A'))
            if visit.get('vitals'):
                v = visit['vitals']
                p3 = doc.add_paragraph(style='List Bullet 2')
                p3.add_run('Vitals: ')
                p3.add_run(f"BP {v.get('bloodPressure', 'N/A')}, HR {v.get('heartRate', 'N/A')}, Temp {v.get('temperature', 'N/A')}")
            p4 = doc.add_paragraph(style='List Bullet 2')
            p4.add_run('Assessment: ').bold = True
            p4.add_run(visit.get('assessment', 'N/A'))
            p5 = doc.add_paragraph(style='List Bullet 2')
            p5.add_run('Plan: ').bold = True
            p5.add_run(visit.get('plan', 'N/A'))

    if patient.get('surgicalHistory'):
        doc.add_heading('Surgical History', level=1)
        doc.add_paragraph(patient.get('surgicalHistory'))

    if patient.get('immunizations'):
        doc.add_heading('Immunizations', level=1)
        doc.add_paragraph(patient.get('immunizations'))

    if patient.get('generalNotes'):
        doc.add_heading('General Notes', level=1)
        doc.add_paragraph(patient.get('generalNotes'))

    _add_footer(doc, 'CONFIDENTIAL MEDICAL RECORD\nFor authorized use only\nThis document contains protected health information')
    return doc
